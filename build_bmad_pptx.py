from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Colours ───────────────────────────────────────────────────────────────────
DARK_NAVY  = RGBColor(0x0A, 0x22, 0x46)
MID_NAVY   = RGBColor(0x1D, 0x35, 0x57)
EVAL_BLUE  = RGBColor(0x00, 0x77, 0xB6)
EVAL_FOOT  = RGBColor(0x02, 0x3E, 0x8A)
GREEN_DARK = RGBColor(0x2D, 0x6A, 0x4F)
GREEN_FOOT = RGBColor(0x38, 0x66, 0x41)
HUMAN_YLW  = RGBColor(0xE9, 0xC4, 0x6A)
HUMAN_TXT  = RGBColor(0x3D, 0x2B, 0x00)
HUMAN_CMD  = RGBColor(0x1A, 0x3A, 0x6A)
CMD_BG     = RGBColor(0xFF, 0xF3, 0xCC)
LIGHT_PNL  = RGBColor(0xF8, 0xFA, 0xFF)
SUB_HDR    = RGBColor(0xE4, 0xED, 0xF8)
OUT_BG     = RGBColor(0xF0, 0xF7, 0xFF)
SEP_LINE   = RGBColor(0x1A, 0x2D, 0x58)
FTR_SEP    = RGBColor(0xC9, 0xD3, 0xDF)
FTR_GRAY   = RGBColor(0x5A, 0x5A, 0x5A)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
OUT_JIRA   = RGBColor(0x00, 0x40, 0xAA)
OUT_GH     = RGBColor(0x1A, 0x52, 0x36)
OUT_FILE   = RGBColor(0x33, 0x40, 0x60)
BODY_TXT   = RGBColor(0x0A, 0x22, 0x46)
LOOP_FG    = RGBColor(0xFF, 0xCC, 0x80)
PLAN_CLR   = RGBColor(0xE6, 0x77, 0x00)
BUILD_CLR  = RGBColor(0x00, 0x77, 0xB6)
SHIP_CLR   = RGBColor(0x2D, 0x6A, 0x4F)
BADGE_BG   = RGBColor(0xFF, 0xF0, 0xAA)
BADGE_TXT  = RGBColor(0x4A, 0x2A, 0x00)
# flow indicator colours
PASS_BG    = RGBColor(0xC8, 0xED, 0xCE)
PASS_BD    = RGBColor(0x28, 0x7A, 0x56)
PASS_TXT   = RGBColor(0x0A, 0x3D, 0x1F)
RETRY_BG   = RGBColor(0xFF, 0xF5, 0xCC)
RETRY_BD   = RGBColor(0xE6, 0x77, 0x00)
RETRY_TXT  = RGBColor(0x5A, 0x30, 0x00)
ESC_H_BG   = RGBColor(0xFF, 0xE8, 0xD6)
ESC_H_BD   = RGBColor(0xC0, 0x20, 0x10)
ESC_H_TXT  = RGBColor(0x7A, 0x10, 0x00)
ESC_A_BG   = RGBColor(0xD6, 0xE8, 0xFF)
ESC_A_BD   = RGBColor(0x00, 0x55, 0x99)
ESC_A_TXT  = RGBColor(0x00, 0x2D, 0x6E)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.50)

# ── Primitives ────────────────────────────────────────────────────────────────
def blank():
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = WHITE
    return sl

def box(sl, x, y, w, h, fill, border=None, bw=Pt(1)):
    s = sl.shapes.add_shape(1, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if border:
        s.line.color.rgb = border; s.line.width = bw
    else:
        s.line.fill.background()
    return s

def lbl(sl, text, x, y, w, h, sz=Pt(10), clr=DARK_NAVY,
        bold=False, align=PP_ALIGN.LEFT, italic=False):
    tb = sl.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p  = tf.paragraphs[0]; p.alignment = align
    r  = p.add_run()
    r.text = text; r.font.size = sz; r.font.color.rgb = clr
    r.font.bold = bold; r.font.italic = italic; r.font.name = "Calibri"

def lbl2(sl, line1, line2, x, y, w, h, sz1=Pt(10), sz2=Pt(9),
         clr=DARK_NAVY, bold1=True, align=PP_ALIGN.CENTER):
    tb = sl.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p1 = tf.paragraphs[0]; p1.alignment = align
    r1 = p1.add_run(); r1.text = line1
    r1.font.size = sz1; r1.font.color.rgb = clr
    r1.font.bold = bold1; r1.font.name = "Calibri"
    p2 = tf.add_paragraph(); p2.alignment = align
    r2 = p2.add_run(); r2.text = line2
    r2.font.size = sz2; r2.font.color.rgb = clr
    r2.font.bold = False; r2.font.name = "Calibri"

def mlbl(sl, lines, clrs, x, y, w, h, sz=Pt(10),
         bold_rows=(), align=PP_ALIGN.LEFT, spacing=Pt(2)):
    tb = sl.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, (line, clr) in enumerate(zip(lines, clrs)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_before = spacing
        r = p.add_run()
        r.text = line; r.font.size = sz; r.font.color.rgb = clr
        r.font.bold = (i in bold_rows); r.font.name = "Calibri"

def arr(sl, x, y, sym="▶", clr=DARK_NAVY, sz=Pt(12)):
    tb = sl.shapes.add_textbox(x, y, Inches(0.26), Inches(0.28))
    tf = tb.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = sym; r.font.size = sz
    r.font.color.rgb = clr; r.font.bold = True; r.font.name = "Calibri"

# ── Column layout ─────────────────────────────────────────────────────────────
HX = Inches(0.15);  HW = Inches(1.58)   # human
AX = Inches(1.88);  AW = Inches(5.55)   # agent
FX = Inches(7.55);  FW = Inches(1.45)   # flow indicator (NEW)
OX = Inches(9.10);  OW = Inches(4.05)   # output

# ── Chrome ────────────────────────────────────────────────────────────────────
def slide_header(sl, title, num, phase_clr):
    box(sl, Inches(0.25), Inches(0.28), Inches(0.13), Inches(0.54), phase_clr)
    lbl(sl, title, Inches(0.50), Inches(0.30), Inches(11.0), Inches(0.52),
        sz=Pt(22), clr=DARK_NAVY, bold=True)
    lbl(sl, num, Inches(11.6), Inches(0.30), Inches(1.5), Inches(0.4),
        sz=Pt(12), clr=FTR_GRAY, align=PP_ALIGN.RIGHT)
    box(sl, Inches(0.25), Inches(0.95), Inches(12.83), Inches(0.013), SEP_LINE)

def slide_footer(sl):
    box(sl, Inches(0.25), Inches(7.13), Inches(12.83), Inches(0.008), FTR_SEP)
    lbl(sl, "Driving Intelligence Faster",
        Inches(0.25), Inches(7.18), Inches(4.5), Inches(0.26),
        sz=Pt(9), clr=DARK_NAVY, bold=True)
    lbl(sl, "Write guard:  JIRA + GitHub writes require explicit human approval",
        Inches(4.8), Inches(7.18), Inches(8.3), Inches(0.26),
        sz=Pt(8.5), clr=FTR_GRAY, align=PP_ALIGN.RIGHT, italic=True)

def col_headers(sl):
    defs = [
        (HX, HW, "HUMAN ROLE"),
        (AX, AW, "AGENT"),
        (FX, FW, "PASS / RETRY / EXHAUST"),
        (OX, OW, "OUTPUT"),
    ]
    for x, w, txt in defs:
        box(sl, x, Inches(1.00), w, Inches(0.28), SUB_HDR)
        lbl(sl, txt, x + Inches(0.08), Inches(1.01), w - Inches(0.1), Inches(0.26),
            sz=Pt(10.5), clr=DARK_NAVY, bold=True)

# ── Human box ─────────────────────────────────────────────────────────────────
def human_box(sl, y, bh, rid, rname, cmd,
              id_sz=Pt(16), nm_sz=Pt(11), cmd_sz=Pt(9.5)):
    box(sl, HX, y, HW, bh, HUMAN_YLW,
        border=RGBColor(0xC8, 0x9A, 0x00), bw=Pt(1.5))
    # role id
    lbl(sl, rid, HX + Inches(0.06), y + Inches(0.10),
        HW - Inches(0.10), Inches(0.32),
        sz=id_sz, clr=HUMAN_TXT, bold=True, align=PP_ALIGN.CENTER)
    # role name
    lbl(sl, rname, HX + Inches(0.06), y + Inches(0.40),
        HW - Inches(0.10), Inches(0.38),
        sz=nm_sz, clr=HUMAN_TXT, align=PP_ALIGN.CENTER)
    # /command strip
    ch = Inches(0.30)
    cy = y + bh - ch - Inches(0.08)
    box(sl, HX + Inches(0.06), cy, HW - Inches(0.10), ch,
        CMD_BG, border=RGBColor(0xC0, 0x90, 0x10), bw=Pt(0.75))
    lbl(sl, cmd, HX + Inches(0.08), cy + Inches(0.04),
        HW - Inches(0.14), ch - Inches(0.06),
        sz=cmd_sz, clr=HUMAN_CMD, bold=True, align=PP_ALIGN.CENTER)

# ── Agent box — header + evaluator + gate only (no body bullets) ───────────────
def agent_box(sl, y, total_h, hdr_h, sub_h, gate_h,
              hdr_clr, gate_clr, name, cmd, evaluator, gate_text, max_att,
              hdr_sz=Pt(13), sub_sz=Pt(11), gate_sz=Pt(11)):
    body_h = total_h - hdr_h - sub_h - gate_h
    box(sl, AX, y, AW, total_h, LIGHT_PNL)

    # ── header ──
    box(sl, AX, y, AW, hdr_h, hdr_clr)
    lbl(sl, f"{name}",
        AX + Inches(0.14), y + Inches(0.05),
        AW - Inches(2.0), hdr_h - Inches(0.06),
        sz=hdr_sz, clr=WHITE, bold=True)
    lbl(sl, cmd,
        AX + Inches(0.14), y + Inches(0.05) + hdr_h * 0.48,
        AW - Inches(2.0), hdr_h * 0.52,
        sz=Pt(hdr_sz.pt * 0.82), clr=RGBColor(0xCC, 0xE5, 0xFF))
    # ↺ retry badge (right side of header)
    bw = Inches(1.72)
    bx = AX + AW - bw - Inches(0.06)
    bh_badge = hdr_h - Inches(0.08)
    box(sl, bx, y + Inches(0.04), bw, bh_badge, BADGE_BG,
        border=RGBColor(0xC8, 0x9A, 0x00), bw=Pt(0.75))
    lbl(sl, f"↺  max {max_att} attempts",
        bx + Inches(0.06), y + Inches(0.06),
        bw - Inches(0.10), bh_badge - Inches(0.04),
        sz=Pt(9.5), clr=BADGE_TXT, bold=True, align=PP_ALIGN.CENTER)

    # ── evaluator ──
    sy = y + hdr_h
    box(sl, AX, sy, AW, sub_h, SUB_HDR)
    lbl(sl, f"Evaluator:  {evaluator}",
        AX + Inches(0.14), sy + Inches(0.04),
        AW - Inches(0.18), sub_h - Inches(0.06),
        sz=sub_sz, clr=DARK_NAVY)

    # ── optional body (blank panel — spacious feel) ──
    if body_h > Inches(0.04):
        box(sl, AX, sy + sub_h, AW, body_h, LIGHT_PNL)

    # ── gate ──
    gy = y + total_h - gate_h
    box(sl, AX, gy, AW, gate_h, gate_clr)
    lbl(sl, gate_text,
        AX + Inches(0.14), gy + Inches(0.04),
        AW - Inches(0.18), gate_h - Inches(0.06),
        sz=gate_sz, clr=WHITE, bold=True)

# ── Flow indicator — 3 coloured sections per row ─────────────────────────────
def flow_indicator(sl, y, row_h, max_att, esc_dest, esc_mode):
    third = row_h / 3.0

    # PASS (top)
    box(sl, FX, y, FW, third, PASS_BG, border=PASS_BD, bw=Pt(1))
    lbl2(sl, "PASS", "▼  to output", FX, y + Inches(0.04),
         FW, third - Inches(0.06),
         sz1=Pt(11), sz2=Pt(9), clr=PASS_TXT, bold1=True)

    # RETRY (middle)
    ry = y + third
    box(sl, FX, ry, FW, third, RETRY_BG, border=RETRY_BD, bw=Pt(1))
    lbl2(sl, "FAIL  ↺", f"retry  ×{max_att}", FX, ry + Inches(0.04),
         FW, third - Inches(0.06),
         sz1=Pt(11), sz2=Pt(9), clr=RETRY_TXT, bold1=True)

    # EXHAUST (bottom)
    ey = ry + third
    eh = row_h - 2 * third
    if esc_mode == "human":
        bg, bd, tc, sym = ESC_H_BG, ESC_H_BD, ESC_H_TXT, "⚠"
    else:
        bg, bd, tc, sym = ESC_A_BG, ESC_A_BD, ESC_A_TXT, "↩"
    box(sl, FX, ey, FW, eh, bg, border=bd, bw=Pt(1))
    lbl2(sl, f"{sym} EXHAUST", esc_dest, FX, ey + Inches(0.04),
         FW, eh - Inches(0.06),
         sz1=Pt(10), sz2=Pt(9), clr=tc, bold1=True)

# ── Output panel ──────────────────────────────────────────────────────────────
def output_panel(sl, y, bh, lines, sz=Pt(11)):
    box(sl, OX, y, OW, bh, OUT_BG)
    ls = [l[0] for l in lines]
    cs = [l[1] for l in lines]
    mlbl(sl, ls, cs,
         OX + Inches(0.16), y + Inches(0.14),
         OW - Inches(0.22), bh - Inches(0.18),
         sz=sz, spacing=Pt(5))

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 1 — PLAN  (3 rows)
# Available y: 1.30 → 7.08 = 5.78 in
# 3 rows  +  2 × (arrow 0.08 + gap 0.05)  →  per row = (5.78 − 0.26) / 3 = 1.84 in
# ─────────────────────────────────────────────────────────────────────────────
P_ROW = Inches(1.84)
P_GAP = Inches(0.13)   # space between rows (includes arrow)
P_Y   = Inches(1.30)
P_HH  = Inches(0.52)
P_SH  = Inches(0.34)
P_GH  = Inches(0.34)

s1 = blank()
slide_header(s1, "BMAD  —  PHASE 1 :  PLAN", "1 of 3", PLAN_CLR)
col_headers(s1)

PLAN_DATA = [
    # (rid, rname, hcmd, name, agent_cmd, evaluator, gate, gate_clr, max_att, output, esc_dest, esc_mode)
    ("H1", "Product\nManager", "/pm-eval-loop\n\"requirement\"",
     "1.  PM EVAL LOOP", "/pm-eval-loop",
     "PRD quality — 10 items",
     "🏷  bmad-prd-approved  on Epic", EVAL_FOOT, 2,
     [("JIRA:  Epic  +  bmad-prd-approved",  OUT_JIRA),
      ("Confluence:  PRD page",               OUT_FILE)],
     "→  HUMAN", "human"),

    ("H2", "Business\nAnalyst", "/ba-eval-loop\nepic:{KEY}",
     "2.  BA EVAL LOOP", "/ba-eval-loop",
     "Story quality — 9 items / story",
     "🏷  bmad-dor-ready  on each Story", EVAL_FOOT, 2,
     [("JIRA:  Stories  +  ACs",             OUT_JIRA),
      ("JIRA:  bmad-dor-ready",              OUT_JIRA)],
     "→  /pm-eval-loop", "agent"),

    ("H3", "Architect", "/arch-eval-loop\nSCRUM-{n}",
     "3.  ARCH EVAL LOOP", "/arch-eval-loop",
     "Definition of Ready — 11 items",
     "🏷  bmad-arch-approved", EVAL_FOOT, 2,
     [("JIRA:  bmad-arch-approved",           OUT_JIRA),
      ("File:  tech design .md",              OUT_FILE)],
     "→  /ba-eval-loop", "agent"),
]

for i, (rid, rname, hcmd, name, cmd, ev, gate, gc, ma, output, esc, em) in enumerate(PLAN_DATA):
    y = P_Y + i * (P_ROW + P_GAP)
    human_box(s1, y + Inches(0.10), P_ROW - Inches(0.14),
              rid, rname, hcmd, id_sz=Pt(16), nm_sz=Pt(11), cmd_sz=Pt(9.5))
    arr(s1, HX + HW + Inches(0.04), y + P_ROW/2 - Inches(0.14), sz=Pt(13))
    agent_box(s1, y, P_ROW, P_HH, P_SH, P_GH, EVAL_BLUE, gc,
              name, cmd, ev, gate, ma,
              hdr_sz=Pt(13), sub_sz=Pt(11), gate_sz=Pt(11))
    flow_indicator(s1, y, P_ROW, ma, esc, em)
    arr(s1, OX - Inches(0.28), y + P_ROW/2 - Inches(0.14), sz=Pt(13))
    output_panel(s1, y + Inches(0.10), P_ROW - Inches(0.14), output, sz=Pt(11))
    if i < 2:
        arr(s1, AX + AW/2 - Inches(0.13),
            y + P_ROW + Inches(0.01), "▼", sz=Pt(11))

lbl(s1, "▼  APPROVED  →  Slide 2 :  BUILD",
    AX + Inches(1.8), P_Y + 3*(P_ROW+P_GAP) - Inches(0.20),
    Inches(3.5), Inches(0.24),
    sz=Pt(10), clr=PLAN_CLR, bold=True, align=PP_ALIGN.CENTER)
slide_footer(s1)

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 2 — BUILD  (2 rows + REQUEST CHANGES loop box)
# Available: 1.56 → 7.08 = 5.52 in
# Loop box 0.48  +  1 gap 0.13  →  per row = (5.52 − 0.48 − 0.13) / 2 = 2.455 in
# ─────────────────────────────────────────────────────────────────────────────
B_ROW = Inches(2.455)
B_GAP = Inches(0.13)
B_LPH = Inches(0.48)
B_Y   = Inches(1.56)
B_HH  = Inches(0.56)
B_SH  = Inches(0.38)
B_GH  = Inches(0.38)

s2 = blank()
slide_header(s2, "BMAD  —  PHASE 2 :  BUILD", "2 of 3", BUILD_CLR)
col_headers(s2)
lbl(s2, "Requires:  bmad-dor-ready  +  bmad-arch-approved     "
        "◄── re-entered from Slide 3 on impl gap",
    AX + Inches(0.12), Inches(1.31), AW, Inches(0.22),
    sz=Pt(9.5), clr=RGBColor(0xC0, 0x20, 0x10), bold=True)

BUILD_DATA = [
    ("H4", "Developer", "/dev-eval-loop\nSCRUM-{n}",
     "4.  DEV EVAL LOOP", "/dev-eval-loop",
     "12-cat inline review  +  mvn test  +  JaCoCo ≥ 80%",
     "🏷  git push  +  PR raised on first passing attempt", EVAL_FOOT, 3,
     [("Files:  source + migrations",          OUT_FILE),
      ("GitHub:  git push + PR raised",        OUT_GH),
      ("CI:  mvn test green · JaCoCo ≥ 80%",  OUT_GH)],
     "→  HUMAN", "human"),

    ("H4", "Developer", "/code-review-eval-loop\nSCRUM-{n}",
     "5.  CODE REVIEW EVAL LOOP", "/code-review-eval-loop",
     "12-category PR audit  ·  round state tracked across invocations",
     "🏷  bmad-code-review-approved  ·  PR  =  APPROVED", GREEN_FOOT, 3,
     [("GitHub:  PR  APPROVED",                OUT_GH),
      ("JIRA:  bmad-code-review-approved",     OUT_JIRA),
      ("JIRA:  REQUEST CHANGES comment",       OUT_JIRA)],
     "→  HUMAN", "human"),
]

for i, (rid, rname, hcmd, name, cmd, ev, gate, gc, ma, output, esc, em) in enumerate(BUILD_DATA):
    y = B_Y + i * (B_ROW + B_GAP)
    if i == 1:
        y += B_LPH
    human_box(s2, y + Inches(0.18), B_ROW - Inches(0.26),
              rid, rname, hcmd, id_sz=Pt(16), nm_sz=Pt(11), cmd_sz=Pt(9.5))
    arr(s2, HX + HW + Inches(0.04), y + B_ROW/2 - Inches(0.14), sz=Pt(13))
    agent_box(s2, y, B_ROW, B_HH, B_SH, B_GH, EVAL_BLUE, gc,
              name, cmd, ev, gate, ma,
              hdr_sz=Pt(14), sub_sz=Pt(12), gate_sz=Pt(12))
    flow_indicator(s2, y, B_ROW, ma, esc, em)
    arr(s2, OX - Inches(0.28), y + B_ROW/2 - Inches(0.14), sz=Pt(13))
    output_panel(s2, y + Inches(0.18), B_ROW - Inches(0.28), output, sz=Pt(12))

# REQUEST CHANGES loop box
loop_y = B_Y + B_ROW + B_GAP
arr(s2, AX + AW/2 - Inches(0.13), loop_y - Inches(0.06), "▼", sz=Pt(11))
box(s2, AX + Inches(0.5), loop_y, AW - Inches(1.0), B_LPH,
    RGBColor(0x1B, 0x3A, 0x5C))
lbl(s2, "REQUEST CHANGES  →  dev fixes on branch  "
        "→  re-run /dev-eval-loop  →  re-run /code-review-eval-loop",
    AX + Inches(0.66), loop_y + Inches(0.10),
    AW - Inches(1.2), B_LPH - Inches(0.10),
    sz=Pt(12), clr=LOOP_FG, bold=True)

lbl(s2, "▼  APPROVED  →  Slide 3 :  SHIP",
    AX + Inches(1.8), B_Y + 2*(B_ROW+B_GAP) + B_LPH - Inches(0.06),
    Inches(3.5), Inches(0.24),
    sz=Pt(10), clr=BUILD_CLR, bold=True, align=PP_ALIGN.CENTER)
slide_footer(s2)

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 3 — SHIP  (4 rows + merge marker)
# Available: 1.54 → 7.08 = 5.54 in
# merge 0.28  +  3 gaps (0.11 each)  →  per row = (5.54 − 0.28 − 0.33) / 4 = 1.2325 in
# ─────────────────────────────────────────────────────────────────────────────
S_ROW = Inches(1.23)
S_GAP = Inches(0.11)
S_MGH = Inches(0.28)
S_Y   = Inches(1.54)
S_HH  = Inches(0.40)
S_SH  = Inches(0.27)
S_GH  = Inches(0.27)

s3 = blank()
slide_header(s3, "BMAD  —  PHASE 3 :  SHIP", "3 of 3", SHIP_CLR)
col_headers(s3)
lbl(s3, "Requires:  bmad-code-review-approved  ·  PR open",
    AX + Inches(0.12), Inches(1.32), AW, Inches(0.20),
    sz=Pt(9.5), clr=DARK_NAVY, italic=True)

# Compute y positions — merge between row 0 (QA) and row 1 (DevOps)
ys = []
y = S_Y
for i in range(4):
    if i == 1:
        y += S_MGH + Inches(0.08)
    ys.append(y)
    y += S_ROW + S_GAP

SHIP_DATA = [
    # (rid, rname, hcmd, name, cmd, evaluator, gate, gate_clr, max_att, output, esc_dest, esc_mode, hdr_clr)
    ("H4", "Developer", "/qa-eval-loop\nSCRUM-{n}",
     "6.  QA EVAL LOOP", "/qa-eval-loop",
     "all ACs + 4 edge cases · Docker stack  (PRE-MERGE)",
     "🏷  JIRA → In Review  ·  curl evidence attached", EVAL_FOOT, 3,
     [("JIRA:  → In Review",              OUT_JIRA),
      ("JIRA:  curl evidence comment",    OUT_JIRA)],
     "→  HUMAN", "human", EVAL_BLUE),

    ("H4", "Developer", "/devops-eval-loop\npr:{n}",
     "8.  DEVOPS EVAL LOOP", "/devops-eval-loop",
     "CI run = green  (POST-MERGE — on pipeline failure)",
     "🏷  CI run = success", GREEN_FOOT, 3,
     [("GitHub:  fix commit pushed",      OUT_GH),
      ("CI:  green",                      OUT_GH)],
     "→  HUMAN", "human", EVAL_BLUE),

    ("H5", "Any Team\nMember", "/sm-eval-loop",
     "7.  SM EVAL LOOP", "/sm-eval-loop",
     "DoD gate 7 items  +  report completeness 8 items",
     "🏷  JIRA → Done  ·  release notes → Confluence", GREEN_DARK, 2,
     [("JIRA:  → Done",                  OUT_JIRA),
      ("Confluence:  release notes",     OUT_FILE)],
     "→  HUMAN", "human", MID_NAVY),

    ("H4", "Developer", "/doc-eval-loop\nSCRUM-{n}",
     "9.  DOC EVAL LOOP", "/doc-eval-loop",
     "doc completeness 8 items  ·  raises a docs PR",
     "🏷  docs branch pushed  ·  docs PR raised", MID_NAVY, 2,
     [("Files:  openapi.yaml + CLAUDE.md", OUT_FILE),
      ("GitHub:  docs PR raised",          OUT_GH)],
     "→  /arch-eval-loop", "agent", MID_NAVY),
]

for i, (rid, rname, hcmd, name, cmd, ev, gate, gc, ma, output, esc, em, hc) in enumerate(SHIP_DATA):
    yi = ys[i]
    human_box(s3, yi + Inches(0.08), S_ROW - Inches(0.10),
              rid, rname, hcmd, id_sz=Pt(14), nm_sz=Pt(10), cmd_sz=Pt(8.5))
    arr(s3, HX + HW + Inches(0.04), yi + S_ROW/2 - Inches(0.14), sz=Pt(12))
    agent_box(s3, yi, S_ROW, S_HH, S_SH, S_GH, hc, gc,
              name, cmd, ev, gate, ma,
              hdr_sz=Pt(12), sub_sz=Pt(10.5), gate_sz=Pt(10.5))
    flow_indicator(s3, yi, S_ROW, ma, esc, em)
    arr(s3, OX - Inches(0.28), yi + S_ROW/2 - Inches(0.14), sz=Pt(12))
    output_panel(s3, yi + Inches(0.08), S_ROW - Inches(0.12), output, sz=Pt(10.5))
    if i < 3:
        arr(s3, AX + AW/2 - Inches(0.13),
            yi + S_ROW + Inches(0.01), "▼", sz=Pt(11))

# Merge marker
merge_y = ys[0] + S_ROW + S_GAP + Inches(0.02)
box(s3, AX + Inches(0.8), merge_y, AW - Inches(1.6), S_MGH, GREEN_DARK)
lbl(s3, "◆  QA PASS  →  H4 MERGES PR to develop  →  triggers /devops-eval-loop on CI failure",
    AX + Inches(0.95), merge_y + Inches(0.05),
    AW - Inches(1.9), S_MGH - Inches(0.08),
    sz=Pt(11), clr=WHITE, bold=True, align=PP_ALIGN.CENTER)

# QA escalation note (3 paths — unique to QA)
lbl(s3, "⚠  impl gap → Slide 2  ·  design gap → Slide 1  ·  3× exhaust → Human",
    AX + Inches(0.14), ys[0] + S_HH + S_SH + Inches(0.06),
    AW - Inches(0.22), Inches(0.22),
    sz=Pt(9.5), clr=RGBColor(0xC0, 0x20, 0x10), bold=True)

slide_footer(s3)

# ── Save ──────────────────────────────────────────────────────────────────────
out = "/Users/umeshkumarkotaladinnevaradaraj/code/broker-dealer/BMAD_3.0.pptx"
prs.save(out)
print(f"Saved: {out}")
